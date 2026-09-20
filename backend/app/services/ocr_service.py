import os
import fitz  # PyMuPDF
import docx
from PIL import Image
import pytesseract
import re

def extract_text_from_file(file_path: str, filename: str) -> tuple[str, str]:
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    
    if ext == ".pdf":
        text = extract_pdf(file_path)
    elif ext == ".docx":
        text = extract_docx(file_path)
    elif ext in [".txt", ".text"]:
        text = extract_txt(file_path)
    elif ext in [".jpg", ".jpeg", ".png"]:
        text = extract_image(file_path)
    else:
        # Fallback text read
        text = extract_txt(file_path)
        
    if not text.strip():
        text = "No readable text extracted. Please ensure the document is clear."
        
    return text, ext

def extract_pdf(file_path: str) -> str:
    text = ""
    extracted_urls = []
    try:
        doc = fitz.open(file_path)
        for page in doc:
            page_text = page.get_text("text")
            if page_text.strip():
                text += page_text + "\n"
            
            # Extract PDF hyperlink annotations
            try:
                links = page.get_links()
                for link in links:
                    if isinstance(link, dict) and "uri" in link and link["uri"]:
                        uri = link["uri"].strip()
                        if uri and uri not in extracted_urls:
                            extracted_urls.append(uri)
            except Exception as le:
                print(f"Error extracting PDF links: {le}")

            # Fallback OCR for scanned pages or sparse text
            if len(page_text.strip()) < 80:
                try:
                    pix = page.get_pixmap()
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_text = pytesseract.image_to_string(img)
                    if ocr_text.strip():
                        text += "\n" + ocr_text + "\n"
                except Exception as oe:
                    print(f"PyMuPDF OCR note: {oe}")

        # If any hyperlinks were extracted from PDF annotations, append them to text
        if extracted_urls:
            text += "\n\nExtracted Contact Links:\n" + "\n".join(extracted_urls) + "\n"

    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text

def extract_docx(file_path: str) -> str:
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"Error reading DOCX, trying raw text fallback: {e}")
        text = extract_txt(file_path)
    return text

def extract_txt(file_path: str) -> str:
    text = ""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    except Exception as e:
        print(f"Error reading text file: {e}")
    return text

def extract_image(file_path: str) -> str:
    text = ""
    try:
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
    except Exception as e:
        print(f"Error reading image OCR: {e}")
    return text

